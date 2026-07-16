from __future__ import annotations

import base64
import hashlib
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document


IMAGE_LIMIT = 10 * 1024 * 1024
DOCX_LIMIT = 20 * 1024 * 1024
DOCX_TEXT_LIMIT = 100_000
ALLOWED_IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _valid_image(extension: str, data: bytes) -> bool:
    if extension == ".png":
        return data.startswith(b"\x89PNG\r\n\x1a\n")
    if extension in {".jpg", ".jpeg"}:
        return data.startswith(b"\xff\xd8")
    if extension == ".webp":
        return len(data) >= 12 and data.startswith(b"RIFF") and data[8:12] == b"WEBP"
    if extension == ".gif":
        return data.startswith((b"GIF87a", b"GIF89a"))
    return False


def extract_docx_text(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as error:
        raise ValueError("DOCX 文件损坏或无法解析") from error
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = "\n".join(blocks).strip()
    if not text:
        raise ValueError("DOCX 中没有可读取的文字")
    if len(text) > DOCX_TEXT_LIMIT:
        text = text[:DOCX_TEXT_LIMIT] + "\n[文档内容已截断]"
    return text


class StudioAssetService:
    def __init__(
        self,
        repository: Any,
        root: str | Path,
        public_prefix: str = "/api/assets",
        generated_root: str | Path | None = None,
    ):
        self.repository = repository
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.generated_root = Path(generated_root) if generated_root else self.root
        self.public_prefix = public_prefix.rstrip("/")

    def _directory(self, user_id: str) -> Path:
        directory = self.root / hashlib.sha256(user_id.encode("utf-8")).hexdigest()[:16]
        directory.mkdir(parents=True, exist_ok=True)
        return directory

    def upload(self, *, user_id: str, filename: str, content_type: str, data: bytes) -> dict[str, Any]:
        safe_name = Path(filename or "").name
        extension = Path(safe_name).suffix.lower()
        extracted_text = None
        if extension in ALLOWED_IMAGE_TYPES:
            kind, mime_type, limit = "image", ALLOWED_IMAGE_TYPES[extension], IMAGE_LIMIT
            if content_type and content_type not in {mime_type, "application/octet-stream"}:
                raise ValueError("图片 MIME 类型与扩展名不一致")
            if not _valid_image(extension, data):
                raise ValueError("图片文件签名无效")
        elif extension == ".docx":
            kind, mime_type, limit = "docx", DOCX_MIME, DOCX_LIMIT
            if content_type and content_type not in {DOCX_MIME, "application/octet-stream"}:
                raise ValueError("DOCX MIME 类型无效")
            if not data.startswith(b"PK"):
                raise ValueError("DOCX 文件签名无效")
            extracted_text = extract_docx_text(data)
        else:
            raise ValueError("不支持该附件格式")
        if not data:
            raise ValueError("附件不能为空")
        if len(data) > limit:
            raise ValueError("附件超过大小限制")
        asset_id = self.repository._id()
        relative = Path(hashlib.sha256(user_id.encode("utf-8")).hexdigest()[:16]) / f"{asset_id}{extension}"
        target = self.root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        try:
            asset = self.repository.create_asset(
                {
                    "id": asset_id,
                    "user_id": user_id,
                    "source": "uploaded",
                    "kind": kind,
                    "filename": safe_name,
                    "mime_type": mime_type,
                    "size": len(data),
                    "storage_path": relative.as_posix(),
                    "extracted_text": extracted_text,
                }
            )
        except Exception:
            target.unlink(missing_ok=True)
            raise
        return asset

    def bind_assets(self, message_id: str, user_id: str, asset_ids: list[str]) -> None:
        self.validate_assets(user_id, asset_ids)
        self.repository.bind_message_assets(message_id, user_id, asset_ids)

    def validate_assets(self, user_id: str, asset_ids: list[str]) -> list[dict[str, Any]]:
        if len(asset_ids) > 8:
            raise ValueError("单条消息最多 8 个附件")
        assets = self.repository.get_assets(asset_ids, user_id)
        if len(assets) != len(asset_ids) or any(item.get("deleted_at") for item in assets):
            raise ValueError("附件不属于当前用户或已删除")
        if sum(int(item["size"]) for item in assets) > 30 * 1024 * 1024:
            raise ValueError("单条消息附件总大小超过 30 MB")
        return assets

    def content_for(self, asset: dict[str, Any]) -> bytes:
        storage_path = asset.get("storage_path")
        if not storage_path:
            raise FileNotFoundError("附件已删除")
        root = self.generated_root if asset.get("source") == "generated" else self.root
        path = (root / storage_path).resolve()
        if root.resolve() not in path.parents:
            raise FileNotFoundError("附件路径无效")
        return path.read_bytes()

    def message_reference_image_data_url(self, message_id: str, user_id: str) -> str | None:
        assets = self.repository.list_message_assets(message_id, user_id)
        image = next(
            (
                item
                for item in assets
                if item.get("kind") == "image" and not item.get("deleted_at")
            ),
            None,
        )
        if not image:
            return None
        encoded = base64.b64encode(self.content_for(image)).decode("ascii")
        return f"data:{image['mime_type']};base64,{encoded}"

    def delete_asset(self, asset_id: str, user_id: str) -> dict[str, Any]:
        asset = self.repository.get_asset(asset_id, user_id)
        if not asset:
            raise KeyError("附件不存在")
        storage_path = asset.get("storage_path")
        if storage_path:
            root = self.generated_root if asset.get("source") == "generated" else self.root
            (root / storage_path).unlink(missing_ok=True)
        return self.repository.tombstone_asset(asset_id, user_id)

    def register_generated(
        self,
        user_id: str,
        image: Any,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        path = self.generated_root / image.relative_path
        asset_metadata = {
            "mode": image.mode,
            "model": image.model,
            "aspect_ratio": image.aspect_ratio,
            "width": image.width,
            "height": image.height,
            "prompt": image.prompt,
        }
        asset_metadata.update(metadata or {})
        return self.repository.create_asset(
            {
                "id": image.id,
                "user_id": user_id,
                "source": "generated",
                "kind": "image",
                "filename": Path(image.relative_path).name,
                "mime_type": image.media_type,
                "size": path.stat().st_size,
                "storage_path": image.relative_path,
                "metadata": asset_metadata,
            }
        )

    def register_generated_video(
        self,
        *,
        user_id: str,
        job_id: str,
        relative_path: str,
        size: int,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        path = self.generated_root / relative_path
        if not path.is_file():
            raise FileNotFoundError("生成视频文件不存在")
        return self.repository.create_asset(
            {
                "id": job_id,
                "user_id": user_id,
                "source": "generated",
                "kind": "video",
                "filename": Path(relative_path).name,
                "mime_type": "video/mp4",
                "size": size,
                "storage_path": relative_path,
                "metadata": metadata or {},
            }
        )
