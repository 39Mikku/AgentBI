import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from AgentBI.src.repositories.sqlite_chat_repository import SqliteChatRepository


PNG = b"\x89PNG\r\n\x1a\n" + b"image-bytes"


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("项目说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "AgentBI"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class StudioAssetServiceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.repository = SqliteChatRepository(":memory:")
        from AgentBI.src.services.studio_asset_service import StudioAssetService

        self.service = StudioAssetService(
            self.repository,
            Path(self.directory.name),
            public_prefix="/api/assets",
        )
        assistant = self.repository.ensure_default_assistant("user-1")
        self.thread = self.repository.create_conversation(
            {
                "user_id": "user-1",
                "title": "附件测试",
                "temperature": 1.0,
                "context_turns": 8,
                "assistant_id": assistant["_id"],
            }
        )
        self.message = self.repository.create_user_message(
            self.thread["_id"], "user-1", "请阅读附件"
        )

    def tearDown(self):
        self.repository.close()
        self.directory.cleanup()

    def test_uploads_image_and_binds_it_to_a_message(self):
        asset = self.service.upload(
            user_id="user-1",
            filename="diagram.png",
            content_type="image/png",
            data=PNG,
        )

        self.service.bind_assets(self.message["_id"], "user-1", [asset["_id"]])
        linked = self.repository.list_message_assets(self.message["_id"], "user-1")

        self.assertEqual(linked[0]["filename"], "diagram.png")
        self.assertEqual(linked[0]["kind"], "image")
        self.assertTrue(Path(self.directory.name, linked[0]["storage_path"]).is_file())
        self.assertEqual(linked[0]["url"], f"/api/assets/{asset['_id']}/content")

    def test_extracts_docx_paragraphs_and_tables_once(self):
        asset = self.service.upload(
            user_id="user-1",
            filename="notes.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=docx_bytes(),
        )

        self.assertEqual(asset["kind"], "docx")
        self.assertIn("项目说明", asset["extracted_text"])
        self.assertIn("名称 | AgentBI", asset["extracted_text"])

    def test_rejects_invalid_and_cross_user_assets(self):
        with self.assertRaisesRegex(ValueError, "不支持"):
            self.service.upload(
                user_id="user-1",
                filename="report.pdf",
                content_type="application/pdf",
                data=b"%PDF",
            )
        asset = self.service.upload(
            user_id="user-1",
            filename="diagram.png",
            content_type="image/png",
            data=PNG,
        )
        with self.assertRaisesRegex(ValueError, "不属于"):
            self.service.bind_assets(self.message["_id"], "user-2", [asset["_id"]])

    def test_delete_tombstones_file_and_all_derived_text(self):
        asset = self.service.upload(
            user_id="user-1",
            filename="diagram.png",
            content_type="image/png",
            data=PNG,
        )
        self.repository.update_asset_derived_text(
            asset["_id"], "user-1", vision_summary="一张架构图"
        )
        self.service.bind_assets(self.message["_id"], "user-1", [asset["_id"]])

        deleted = self.service.delete_asset(asset["_id"], "user-1")
        linked = self.repository.list_message_assets(self.message["_id"], "user-1")[0]

        self.assertIsNotNone(deleted["deleted_at"])
        self.assertIsNone(linked["storage_path"])
        self.assertIsNone(linked["vision_summary"])
        self.assertIsNone(linked["extracted_text"])
        self.assertEqual(linked["status"], "deleted")

    def test_branch_reuses_asset_and_preserves_message_link(self):
        asset = self.service.upload(
            user_id="user-1",
            filename="diagram.png",
            content_type="image/png",
            data=PNG,
        )
        self.service.bind_assets(self.message["_id"], "user-1", [asset["_id"]])

        branch = self.repository.create_branch_conversation(
            self.thread["_id"], "user-1", self.message["_id"], "附件分支"
        )
        branch_messages = self.repository.list_messages(branch["_id"], "user-1")

        self.assertEqual(branch_messages[0]["assets"][0]["_id"], asset["_id"])
        self.assertEqual(
            self.repository.get_asset(asset["_id"], "user-1")["storage_path"],
            asset["storage_path"],
        )


if __name__ == "__main__":
    unittest.main()
